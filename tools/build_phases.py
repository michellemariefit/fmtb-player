"""
Build data/fmtb-phase-N.json from the Word workout files.

    pip install python-docx pillow
    python tools/build_phases.py 3 4 5 6 7 8 9 10

Put the Word files in tools/docx/ (ignored by git) named gym-03.docx and
home-03.docx, and so on. Each file holds one table per day (or one table with
"Day N" rows inside it) with columns EXERCISE, REPS, SETS, REST, TEMPO, often
two exercises side by side. Exercise cells carry a YouTube link and a thumbnail.

What it writes:
  data/fmtb-phase-N.json     same shape as phases 1 and 2
  images/<slug>.webp         the Word thumbnail, used until start/end photos exist
Then run tools/grab_frames.py to add the start and end photos.
"""
import io, json, os, re, sys
import docx
from docx.oxml.ns import qn
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "tools", "docx")
IMAGES = os.path.join(ROOT, "images")
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# From the GHL lessons (see README).
META = {
    3: ("Toned Up Prego", "ehr6YV-cIrA"), 4: ("The Lean Mean Baby Making Machine", "UDQ7Rc7ps-E"),
    5: ("Strong Bod Champ", "jO-BLhLEL9A"), 6: ("Leveled Up Prego", "kO5UMYU2OW0"),
    7: ("Fit Mama Blast", "9pFmm1rq45g"), 8: ("The Lean Mama Slay", "-lFGJfiaMXs"),
    9: ("The Strong Mama Blast", "ipKjuf0IvXs"), 10: ("The Chiseled Mom Workout", "4DPwbPldYew"),
}
# The Word files link one YouTube video for two different exercises in a few places.
VIDEO_FIX = {("0WtcBgPcHYs", "overhead"): "uicoxUY-unM"}
# Videos Michelle supplied later for exercises the Word files had no link for, by name.
VIDEO_BY_NAME = [(r"tubing fly", "2Y0TlHw_BhU"), (r"barbell curl", "DfrY9Mgnq58")]
# Exercises swapped after the fact: (old video id) -> (new name, new video id).
SWAP = {"03yTjkLx8Bk": ("Cable Flys", "wA44NHH3its")}
NAME_FIX = [(r"extensi[oó]n", "extension"), (r"[’']s\b", "s"), (r"\s+", " "), (r"\s*\.\s*$", "")]


def yt(url):
    m = re.search(r"(?:youtu\.be/|v=|embed/)([A-Za-z0-9_-]{11})", url or "")
    return m.group(1) if m else None


def slug(s):
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", s.lower())).strip("-")[:60]


def cell_extras(cell, part):
    links, imgs = [], []
    for h in cell._tc.iter(qn("w:hyperlink")):
        rid = h.get(qn("r:id"))
        if rid in part.rels:
            links.append(part.rels[rid].target_ref)
    for b in cell._tc.iter(A_NS + "blip"):
        rid = b.get(qn("r:embed"))
        if rid in part.rels:
            imgs.append(part.rels[rid].target_part)
    return links, imgs


def parse_doc(path):
    """-> {day: [exercise dicts in reading order]}"""
    d = docx.Document(path)
    part = d.part
    days, day = {}, 0
    for t in d.tables:
        for r in t.rows:
            cells, seen = [], set()
            for c in r.cells:  # merged cells repeat; keep each once
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc)); cells.append(c)
            text = " ".join(c.text for c in cells)
            m = re.search(r"\bDay\s*(\d+)", text, re.I)
            if m and len([c for c in cells if c.text.strip()]) <= 2:
                day = int(m.group(1)); continue
            if re.search(r"\bEXERCISE\b", text) and re.search(r"\bREPS\b", text):
                continue
            i = 0
            while i + 4 < len(cells):
                name = re.sub(r"\s+", " ", cells[i].text.replace("\n", " ")).strip(" .")
                vals = [cells[i + k].text.strip() for k in range(1, 5)]
                if name and re.search(r"\d", vals[0] + vals[1]):
                    m = re.match(r"^([A-Z]\d)\.?\s*(.*)$", name)
                    links, imgs = cell_extras(cells[i], part)
                    days.setdefault(day, []).append({
                        "code": m.group(1) if m else None, "name": (m.group(2) if m else name).strip(),
                        "reps": vals[0], "sets": vals[1], "rest": vals[2], "tempo": vals[3],
                        "video": next((yt(l) for l in links if yt(l)), None), "imgs": imgs})
                    i += 5
                else:
                    i += 1
    return days


def clean_name(n):
    for pat, rep in NAME_FIX:
        n = re.sub(pat, rep, n, flags=re.I)
    return n.strip()


def parse_reps(s):
    s = re.sub(r"\s*[\n/]\s*", "-", s.strip())
    s = re.sub(r"\s*-\s*", "-", s)
    s = s.replace(" p side", " per side")
    return int(s) if re.fullmatch(r"\d+", s) else s


def parse_rest(s):
    m = re.search(r"\d+", s or "")
    return int(m.group()) if m else 0


def parse_sets(s):
    m = re.search(r"\d+", s or "")
    return int(m.group()) if m else 1


def note_for(name, reps):
    n = name.lower()
    if "drop set" in n or (isinstance(reps, str) and re.fullmatch(r"\d+-\d+-\d+", reps)):
        return "Drop set: do the first number, lower the weight, do the next, lower again, do the last. No rest between."
    if "1 1/2" in n:
        return "1 1/2 reps: a full rep, then a half rep. That counts as one."
    return ""


def save_thumb(img_part, name):
    path = os.path.join(IMAGES, slug(name) + ".webp")
    if not os.path.exists(path):
        im = Image.open(io.BytesIO(img_part.blob)).convert("RGB")
        im.thumbnail((720, 720))
        im.save(path, "WEBP", quality=82)
    return "images/" + os.path.basename(path)


CORE = re.compile(r"plank|\bcat\b|up up down|crunch|twist|side kick|bird dog", re.I)


def assign_codes(exs):
    """Fill in codes the Word file left blank. Core finishers (Plank, CAT) become their own
    group after the last lettered one; anything else continues the previous code (B2 -> B3)."""
    last, core = None, []
    for e in exs:
        if e["code"]:
            last = e["code"]
        elif CORE.search(e["name"]):
            core.append(e)
        elif last:
            last = last[0] + str(int(last[1:]) + 1); e["code"] = last
        else:
            last = e["code"] = "A1"
    if core:
        letter = chr(ord(max(e["code"][0] for e in exs if e["code"])) + 1) if any(e["code"] for e in exs) else "A"
        for i, e in enumerate(core, 1):
            e["code"] = f"{letter}{i}"
    for e in exs:
        if e.get("code") and e["code"] in [x["code"] for x in exs if x is not e]:
            print("   duplicate code", e["code"], e["name"])
    return exs


def build_version(exs):
    groups = {}
    for e in assign_codes(exs):
        name = clean_name(e["name"])
        reps = parse_reps(e["reps"])
        video = e["video"]
        for (vid, word), fixed in VIDEO_FIX.items():
            if video == vid and word in name.lower():
                video = fixed
        if not video:
            video = next((v for pat, v in VIDEO_BY_NAME if re.search(pat, name, re.I)), None)
        if video in SWAP:
            suffix = re.sub(r"(?i)(cable incline|incline cable) press(es)?", "", name).strip()
            name, video = (SWAP[video][0] + (" " + suffix if suffix else "")), SWAP[video][1]
        ex = {"code": e["code"], "name": name, "reps": reps, "sets": parse_sets(e["sets"]),
              "restSeconds": parse_rest(e["rest"]), "tempo": e["tempo"].strip(" -") , "video": video,
              "photo": save_thumb(e["imgs"][0], name) if e["imgs"] else "", "note": note_for(name, reps)}
        groups.setdefault(e["code"][0], []).append(ex)
    out = []
    for letter in sorted(groups):
        exs = sorted(groups[letter], key=lambda x: int(x["code"][1:]))
        out.append({"label": letter, "exercises": exs})
    return out


def intro_for(n, days, paras):
    allex = [e for d in days for v in d["versions"].values() for g in v for e in g["exercises"]]
    sets = max(set(e["sets"] for e in allex), key=lambda s: sum(e["sets"] == s for e in allex))
    tempos = sorted({e["tempo"] for e in allex if e["tempo"]})
    sizes = {len(g["exercises"]) for d in days for v in d["versions"].values() for g in v}
    kind = "superset" if sizes == {2} else "giant set" if sizes <= {2, 3} else "circuit"
    text = f"{sets} sets. Each lettered group is a {kind}: do every exercise in it back to back, then rest."
    extra = [p for p in paras if len(p) > 60]
    if extra:
        text += " " + extra[0].strip("*").strip()
    return text


def build(n):
    days, paras = {}, []
    for ver in ("gym", "home"):
        path = os.path.join(DOCX, f"{ver}-{n:02d}.docx")
        d = docx.Document(path)
        paras += [p.text.strip() for p in d.paragraphs if p.text.strip()]
        for day, exs in parse_doc(path).items():
            days.setdefault(day, {})[ver] = build_version(exs)
    out_days = [{"day": day, "versions": days[day]} for day in sorted(days)]
    for d in out_days:
        for ver in ("gym", "home"):
            d["versions"].setdefault(ver, d["versions"].get("gym" if ver == "home" else "home", []))
    name, vid = META[n]
    data = {"program": "FMTB", "phase": n, "name": name, "intro": intro_for(n, out_days, dict.fromkeys(paras)),
            "instructionVideo": vid, "days": out_days}
    path = os.path.join(ROOT, "data", f"fmtb-phase-{n}.json")
    json.dump(data, open(path, "w"), indent=2, ensure_ascii=False)
    count = sum(len(g["exercises"]) for d in out_days for v in d["versions"].values() for g in v)
    print(f"Phase {n}: {len(out_days)} days, {count} exercise slots -> {os.path.relpath(path, ROOT)}")


if __name__ == "__main__":
    for arg in sys.argv[1:] or range(3, 11):
        build(int(arg))
